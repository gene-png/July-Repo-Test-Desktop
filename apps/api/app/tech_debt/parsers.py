"""Inventory file parsers - CSV and XLSX.

Master Spec §15 Phase 3: "Capability list ingest (Excel upload + AI
extraction with redaction)". This module turns the raw artifact bytes
into a row-shaped representation the LLM can reason about. The LLM does
the column-mapping; we just give it well-shaped rows.

Phase 3 only supports CSV + XLSX. PDF ingest is a Phase 6 hardening
target (table extraction is a different problem and the inventory
documents Eugene's customers ship are reliably tabular).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from typing import Any


class UnsupportedInventoryFormat(ValueError):
    """Raised when an artifact's MIME isn't a recognized inventory format."""


class EmptyInventoryError(ValueError):
    """Raised when a recognized inventory file yields zero data rows."""


class CorruptInventoryError(ValueError):
    """Raised when a file claims to be XLSX but can't be opened as one.

    A legacy OLE2 .xls mislabeled as XLSX, or a truncated/corrupt upload,
    lands here so the route can return a clean 422 instead of a 500.
    """


# MIME types that the ingest endpoint accepts.
# Legacy application/vnd.ms-excel (.xls, OLE2) is intentionally absent:
# openpyxl cannot read OLE2, so it must be rejected up front (see the route)
# rather than exploding at parse time.
SUPPORTED_MIME = {
    "text/csv": "csv",
    "text/plain": "csv",  # treat .txt as CSV; most inventory exports save this way
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
}

# Max rows we ship to the LLM. Above this and we'd either bust the context
# window or pay a fortune in tokens. v1 inventories are typically 50-300 rows.
MAX_ROWS = 500


@dataclass(frozen=True)
class ParseReport:
    """Metadata about a parse pass, surfaced back to the admin (C-3).

    ``sheet_used`` is the worksheet the rows were pulled from (None for CSV,
    which has no sheets). ``rows_parsed`` is the count of data rows returned
    (post-truncation); ``rows_skipped`` counts blank rows dropped during the
    parse; ``truncated`` is True when the source exceeded MAX_ROWS and the
    tail was dropped before the model saw it.
    """

    sheet_used: str | None
    rows_parsed: int
    rows_skipped: int
    truncated: bool


def kind_for_mime(mime_type: str) -> str:
    try:
        return SUPPORTED_MIME[mime_type]
    except KeyError as exc:
        raise UnsupportedInventoryFormat(
            f"Inventory format {mime_type!r} not supported. Accept CSV or XLSX."
        ) from exc


def parse_inventory(data: bytes, mime_type: str) -> tuple[list[dict[str, Any]], ParseReport]:
    """Parse `data` into row-dicts plus a :class:`ParseReport`.

    The header row becomes the dict keys. Duplicate header labels are
    uniquified (``name`` / ``name_2``); blank header cells and overflow cells
    (cells past the end of the header row) get generated ``col_N`` keys.

    Returns at most MAX_ROWS rows; ``ParseReport.truncated`` is True when the
    input was longer and the tail was dropped.
    """
    kind = kind_for_mime(mime_type)
    if kind == "csv":
        rows, skipped = _parse_csv(data)
        sheet_used: str | None = None
    elif kind == "xlsx":
        rows, sheet_used, skipped = _parse_xlsx(data)
    else:
        raise UnsupportedInventoryFormat(f"Unknown internal kind {kind!r}.")

    truncated = len(rows) > MAX_ROWS
    if truncated:
        rows = rows[:MAX_ROWS]
    report = ParseReport(
        sheet_used=sheet_used,
        rows_parsed=len(rows),
        rows_skipped=skipped,
        truncated=truncated,
    )
    return rows, report


def _normalize_headers(raw_header: list[Any]) -> list[str]:
    """Turn a raw header row into stable, unique string keys.

    Blank cells become ``col_N`` (1-based column position); duplicate labels
    are suffixed (``name`` -> ``name``, ``name_2``, ``name_3``).
    """
    base: list[str] = []
    for i, h in enumerate(raw_header):
        s = "" if h is None else str(h).strip()
        base.append(s if s else f"col_{i + 1}")
    seen: dict[str, int] = {}
    out: list[str] = []
    for name in base:
        if name in seen:
            seen[name] += 1
            out.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 1
            out.append(name)
    return out


def _build_row(headers: list[str], cells: list[Any]) -> dict[str, Any]:
    """Map a data row's cells onto the normalized headers.

    Overflow cells (past the header length) are kept under generated
    ``col_N`` keys so no data is silently dropped.
    """
    out: dict[str, Any] = {}
    for i, v in enumerate(cells):
        key = headers[i] if i < len(headers) else f"col_{i + 1}"
        out[key] = "" if v is None else str(v).strip()
    return out


def _is_blank_row(cells: list[Any]) -> bool:
    return not cells or all(v is None or str(v).strip() == "" for v in cells)


def _parse_csv(data: bytes) -> tuple[list[dict[str, Any]], int]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    try:
        raw_header = next(reader)
    except StopIteration:
        return [], 0
    headers = _normalize_headers(list(raw_header))
    rows: list[dict[str, Any]] = []
    skipped = 0
    for cells in reader:
        if _is_blank_row(list(cells)):
            skipped += 1
            continue
        rows.append(_build_row(headers, list(cells)))
    return rows, skipped


def _parse_xlsx(data: bytes) -> tuple[list[dict[str, Any]], str | None, int]:
    # openpyxl is lazy-imported so test runs that don't touch XLSX don't
    # pay the import cost.
    from zipfile import BadZipFile

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException) as exc:
        raise CorruptInventoryError(
            "This file could not be opened as an .xlsx workbook. "
            "It may be a legacy .xls or a corrupt upload; re-save it as .xlsx "
            "and try again."
        ) from exc

    # Parse every sheet and pick the best candidate: the one with the most
    # non-empty data rows under a non-empty header row (C-3). A workbook whose
    # first sheet is a cover page and whose data lives on sheet 2 must still
    # ingest.
    best: tuple[str, list[str], list[list[Any]], int] | None = None
    for ws in wb.worksheets:
        header, data_rows, skipped = _read_sheet(ws)
        if header is None:
            continue
        if best is None or len(data_rows) > len(best[2]):
            best = (ws.title, header, data_rows, skipped)

    if best is None:
        return [], None, 0

    sheet_title, headers, data_rows, skipped = best
    rows = [_build_row(headers, cells) for cells in data_rows]
    return rows, sheet_title, skipped


def _read_sheet(ws: Any) -> tuple[list[str] | None, list[list[Any]], int]:
    """Read one worksheet: find the header row, collect data rows.

    The header is the first non-empty row (leading blank rows are ignored);
    returns ``(None, [], 0)`` for an entirely empty sheet.
    """
    rows_iter = ws.iter_rows(values_only=True)
    header: list[str] | None = None
    for raw in rows_iter:
        cells = list(raw) if raw is not None else []
        if not _is_blank_row(cells):
            header = _normalize_headers(cells)
            break
    if header is None:
        return None, [], 0

    data_rows: list[list[Any]] = []
    skipped = 0
    for raw in rows_iter:
        cells = list(raw) if raw is not None else []
        if _is_blank_row(cells):
            skipped += 1
            continue
        data_rows.append(cells)
    return header, data_rows, skipped
