"""Task S1-B / B-2 regression: CSF finalize resolves the deliverable gap target.

Before B-2 the CSF finalize path called ``analyze_gaps`` with no target argument,
so the engine default tier (T3) was always used. This drives the HTTP surface end
to end: it sets an engagement target of tier 4 (via the originating
ServiceRequest), scores subcategories so a bounded set falls short, finalizes,
parses the Gap Plan sheet out of the rendered XLSX, and asserts the row count
matches the dashboard endpoint's total_gap_count for the same target. On the
pre-fix default of 3 the tier-3 subcategories would not register as gaps.
"""

from __future__ import annotations

import io
import os
from collections.abc import Iterator
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from app.storage.local import LocalFilesystemStorage
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from tests.conftest import register_admin


@pytest.fixture()
def app_ctx(tmp_path) -> Iterator[tuple[TestClient, sessionmaker, str]]:
    url = f"sqlite:///{tmp_path / 'shield-s1b-csf.db'}"
    os.environ["DATABASE_URL"] = url
    api_root = Path(__file__).resolve().parents[2]
    cfg = Config(str(api_root / "alembic.ini"))
    cfg.set_main_option("script_location", str(api_root / "alembic"))
    cfg.set_main_option("sqlalchemy.url", url)
    command.upgrade(cfg, "head")

    engine = create_engine(url, future=True)
    TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
    storage = LocalFilesystemStorage(tmp_path / "storage")

    from app.db.session import get_db
    from app.main import create_app
    from app.models.client import Client as _Client
    from app.models.client_domain import ClientDomain as _ClientDomain
    from app.routes.artifacts import _storage_dep

    def override_get_db() -> Iterator[Session]:
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    app = create_app()
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[_storage_dep] = lambda: storage

    _seed = TestSession()
    tenant = _Client(legal_name="Test Tenant")
    _seed.add(tenant)
    _seed.flush()
    _seed.add(_ClientDomain(client_id=tenant.id, domain="example.com"))
    _seed.commit()
    cid = str(tenant.id)
    _seed.close()

    with TestClient(app, headers={"X-Client-Id": cid}) as c:
        yield c, TestSession, cid


def _link_engagement_target(
    TestSession: sessionmaker, svc_id: str, cid: str, requested_by: str, tier: int
) -> None:
    """Attach a ServiceRequest carrying the client's CSF target tier and point
    the Service's source_request_id at it, mirroring an intake-opened engagement."""
    import uuid

    from app.models.service import Service
    from app.models.service_request import ServiceRequest, ServiceType

    db = TestSession()
    try:
        sr = ServiceRequest(
            service_type=ServiceType.NIST_CSF,
            client_id=uuid.UUID(cid),
            requested_by=uuid.UUID(requested_by),
            csf_target_tier=tier,
        )
        db.add(sr)
        db.flush()
        svc = db.get(Service, uuid.UUID(svc_id))
        svc.source_request_id = sr.id
        db.commit()
    finally:
        db.close()


@pytest.mark.unit
def test_finalize_gap_target_matches_dashboard(app_ctx) -> None:
    c, TestSession, cid = app_ctx
    admin = register_admin(c, "admin@example.com")
    h = {"Authorization": f"Bearer {admin['tokens']['access_token']}"}

    svc_id = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "Atlas"}).json()[
        "id"
    ]
    _link_engagement_target(TestSession, svc_id, cid, requested_by=admin["user"]["id"], tier=4)

    assessment = c.post(f"/csf/services/{svc_id}/assessments", headers=h).json()
    answers = assessment["answers"]
    # A bounded set falls short (tier 3 < target 4); the rest meet the target
    # (tier 4). Bounded so the count stays under the Gap Plan cap (top_n).
    gap_count_expected = 12
    for i, ans in enumerate(answers):
        tier = 3 if i < gap_count_expected else 4
        c.patch(f"/csf/answers/{ans['id']}", headers=h, json={"maturity_tier": tier})

    c.post(f"/csf/assessments/{assessment['id']}/approve", headers=h)

    fin = c.post(f"/csf/services/{svc_id}/deliverables/finalize", headers=h)
    assert fin.status_code == 201, fin.text
    deliv = fin.json()
    # The resolved engagement target (T4) is printed in the summary line.
    assert "T4" in deliv["summary"], deliv["summary"]

    dash = c.get(f"/csf/services/{svc_id}/gap-analysis?target_tier=4", headers=h)
    assert dash.status_code == 200, dash.text
    total_gap_count = dash.json()["total_gap_count"]
    assert total_gap_count == gap_count_expected

    dl = c.get(
        f"/artifacts/{deliv['xlsx_artifact_id']}/download",
        headers={**h, "X-Client-Id": cid},
    )
    assert dl.status_code == 200
    wb = load_workbook(io.BytesIO(dl.content))
    ws = wb["Gap Plan"]
    data_rows = [
        r for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)
    ]
    assert len(data_rows) == total_gap_count
    assert gap_count_expected <= 20  # guard: sheet caps at top_n, keep comparison exact
