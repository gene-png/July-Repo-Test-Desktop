"""Task S3-A / H-8: CSF action-plan items — routes + Action Plan XLSX sheet."""

from __future__ import annotations

import io
import os
from collections.abc import Iterator
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from tests.conftest import register_admin_resp


@pytest.fixture()
def app_client(tmp_path) -> Iterator[tuple[TestClient, str]]:
    url = f"sqlite:///{tmp_path / 'shield-csf-action.db'}"
    os.environ["DATABASE_URL"] = url
    api_root = Path(__file__).resolve().parents[2]
    cfg = Config(str(api_root / "alembic.ini"))
    cfg.set_main_option("script_location", str(api_root / "alembic"))
    cfg.set_main_option("sqlalchemy.url", url)
    command.upgrade(cfg, "head")
    engine = create_engine(url, future=True)
    TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    from app.db.session import get_db
    from app.main import create_app
    from app.models.client import Client as _Client
    from app.models.client_domain import ClientDomain as _ClientDomain
    from app.routes.artifacts import _storage_dep
    from app.storage.local import LocalFilesystemStorage

    def override_get_db() -> Iterator[Session]:
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    app = create_app()
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[_storage_dep] = lambda: LocalFilesystemStorage(tmp_path / "storage")
    _seed = TestSession()
    tenant = _Client(legal_name="Test Tenant")
    _seed.add(tenant)
    _seed.flush()
    _seed.add(_ClientDomain(client_id=tenant.id, domain="example.com"))
    _seed.commit()
    cid = str(tenant.id)
    with TestClient(app, headers={"X-Client-Id": cid}) as c:
        yield c, cid


def _admin_headers(c: TestClient) -> dict[str, str]:
    r = register_admin_resp(c, "admin@example.com")
    return {"Authorization": f"Bearer {r.json()['tokens']['access_token']}"}


def _valid_code() -> str:
    from app.csf.catalog import all_codes

    return sorted(all_codes())[0]


@pytest.mark.unit
def test_action_item_crud_and_validation(app_client) -> None:
    c, _cid = app_client
    h = _admin_headers(c)
    svc_id = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "CSF"}).json()[
        "id"
    ]
    assess_id = c.post(f"/csf/services/{svc_id}/assessments", headers=h).json()["id"]

    code = _valid_code()
    # Create.
    created = c.post(
        f"/csf/assessments/{assess_id}/action-items",
        headers=h,
        json={"subcategory_code": code, "owner": "Alice", "milestone": "Q3"},
    )
    assert created.status_code == 201, created.text
    item = created.json()
    assert item["status"] == "open"
    assert item["owner"] == "Alice"
    item_id = item["id"]

    # Bad subcategory rejected.
    bad = c.post(
        f"/csf/assessments/{assess_id}/action-items",
        headers=h,
        json={"subcategory_code": "ZZ.NOPE-99"},
    )
    assert bad.status_code == 422, bad.text

    # List.
    listed = c.get(f"/csf/assessments/{assess_id}/action-items", headers=h)
    assert listed.status_code == 200
    assert [i["id"] for i in listed.json()] == [item_id]

    # Patch status.
    patched = c.patch(
        f"/csf/action-items/{item_id}",
        headers=h,
        json={"status": "done", "due_date": "2026-09-01"},
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["status"] == "done"
    assert patched.json()["due_date"] == "2026-09-01"

    # Delete.
    assert c.delete(f"/csf/action-items/{item_id}", headers=h).status_code == 204
    assert c.get(f"/csf/assessments/{assess_id}/action-items", headers=h).json() == []


@pytest.mark.unit
def test_action_item_requires_admin(app_client) -> None:
    c, _cid = app_client
    h = _admin_headers(c)
    svc_id = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "CSF"}).json()[
        "id"
    ]
    assess_id = c.post(f"/csf/services/{svc_id}/assessments", headers=h).json()["id"]
    # No auth header -> not permitted.
    r = c.post(
        f"/csf/assessments/{assess_id}/action-items",
        json={"subcategory_code": _valid_code()},
    )
    assert r.status_code in (401, 403)


@pytest.mark.unit
def test_playbook_xlsx_action_plan_sheet_mirrors_items(app_client) -> None:
    c, cid = app_client
    h = _admin_headers(c)
    svc_id = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "CSF"}).json()[
        "id"
    ]
    assess_id = c.post(f"/csf/services/{svc_id}/assessments", headers=h).json()["id"]

    c.post(f"/csf/services/{svc_id}/profiles/seed", headers=h, json={"tiers": ["high"]})
    rows = c.get(f"/csf/services/{svc_id}/profile/high", headers=h).json()["rows"]
    for row in rows:
        c.patch(
            f"/csf/dimension-scores/{row['id']}",
            headers=h,
            json={"governance": 2, "policy": 2, "target_level": 4},
        )
    code = _valid_code()
    c.post(
        f"/csf/assessments/{assess_id}/action-items",
        headers=h,
        json={"subcategory_code": code, "owner": "Bob", "milestone": "Harden IAM"},
    )
    c.post(f"/csf/assessments/{assess_id}/approve", headers=h)

    ex = c.post(f"/csf/services/{svc_id}/playbook/export", headers=h)
    assert ex.status_code == 200, ex.text
    art = {a["kind"]: a for a in ex.json()["artifacts"]}["xlsx"]

    dh = {**h, "X-Client-Id": cid}
    dl = c.get(f"/artifacts/{art['artifact_id']}/download", headers=dh)
    assert dl.status_code == 200
    wb = load_workbook(io.BytesIO(dl.content))
    assert "Action Plan" in wb.sheetnames
    ws = wb["Action Plan"]
    flat = [str(cell.value) for r in ws.iter_rows() for cell in r if cell.value is not None]
    assert "Bob" in flat
    assert "Harden IAM" in flat
    assert code in flat


@pytest.mark.unit
def test_playbook_xlsx_action_plan_present_when_empty(app_client) -> None:
    c, cid = app_client
    h = _admin_headers(c)
    svc_id = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "CSF"}).json()[
        "id"
    ]
    assess_id = c.post(f"/csf/services/{svc_id}/assessments", headers=h).json()["id"]
    c.post(f"/csf/services/{svc_id}/profiles/seed", headers=h, json={"tiers": ["high"]})
    rows = c.get(f"/csf/services/{svc_id}/profile/high", headers=h).json()["rows"]
    for row in rows:
        c.patch(
            f"/csf/dimension-scores/{row['id']}",
            headers=h,
            json={"governance": 2, "policy": 2, "target_level": 4},
        )
    c.post(f"/csf/assessments/{assess_id}/approve", headers=h)
    ex = c.post(f"/csf/services/{svc_id}/playbook/export", headers=h)
    assert ex.status_code == 200, ex.text
    art = {a["kind"]: a for a in ex.json()["artifacts"]}["xlsx"]
    dh = {**h, "X-Client-Id": cid}
    dl = c.get(f"/artifacts/{art['artifact_id']}/download", headers=dh)
    wb = load_workbook(io.BytesIO(dl.content))
    ws = wb["Action Plan"]
    flat = [str(cell.value) for r in ws.iter_rows() for cell in r if cell.value is not None]
    assert "No action items yet" in flat
