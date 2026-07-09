"""Task S1-B / B-3: the CSF playbook export scoring gate + Unscored fallback.

The playbook export must refuse to render unless every in-scope dimension-score
row has been scored (scored_at stamped) AND the assessment is approved. Seeding
alone leaves rows unscored; the export then 409s with a count. Once every row is
scored (via PATCH here) and the assessment approved, the export succeeds and no
cell reads "Unscored". A scored-but-unapproved assessment is still refused.
"""

from __future__ import annotations

import io
import os
from collections.abc import Iterator
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from app.storage.local import LocalFilesystemStorage
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from tests.conftest import register_admin_resp


@pytest.fixture()
def app_ctx(tmp_path) -> Iterator[tuple[TestClient, str]]:
    url = f"sqlite:///{tmp_path / 'shield-s1b-gate.db'}"
    os.environ["DATABASE_URL"] = url
    api_root = Path(__file__).resolve().parents[2]
    cfg = Config(str(api_root / "alembic.ini"))
    cfg.set_main_option("script_location", str(api_root / "alembic"))
    cfg.set_main_option("sqlalchemy.url", url)
    command.upgrade(cfg, "head")

    engine = create_engine(url, future=True)
    TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
    storage = LocalFilesystemStorage(tmp_path / "storage")

    from app.db.session import get_db
    from app.main import create_app
    from app.models.client import Client as _Client
    from app.models.client_domain import ClientDomain as _ClientDomain
    from app.routes.artifacts import _storage_dep

    def override_get_db() -> Iterator[Session]:
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    app = create_app()
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[_storage_dep] = lambda: storage

    _seed = TestSession()
    tenant = _Client(legal_name="Test Tenant")
    _seed.add(tenant)
    _seed.flush()
    _seed.add(_ClientDomain(client_id=tenant.id, domain="example.com"))
    _seed.commit()
    cid = str(tenant.id)
    _seed.close()

    with TestClient(app, headers={"X-Client-Id": cid}) as c:
        yield c, cid


def _bootstrap(c: TestClient) -> tuple[dict, str, str]:
    r = register_admin_resp(c, "admin@example.com")
    h = {"Authorization": f"Bearer {r.json()['tokens']['access_token']}"}
    svc_id = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "CSF"}).json()[
        "id"
    ]
    assess_id = c.post(f"/csf/services/{svc_id}/assessments", headers=h).json()["id"]
    c.post(f"/csf/services/{svc_id}/profiles/seed", headers=h, json={"tiers": ["high"]})
    return h, svc_id, assess_id


def _score_all(c: TestClient, h: dict, svc_id: str) -> None:
    rows = c.get(f"/csf/services/{svc_id}/profile/high", headers=h).json()["rows"]
    for row in rows:
        c.patch(
            f"/csf/dimension-scores/{row['id']}",
            headers=h,
            json={"governance": 2, "policy": 2, "implementation": 1},
        )


@pytest.mark.unit
def test_export_gate_blocks_unscored_rows(app_ctx) -> None:
    c, _cid = app_ctx
    h, svc_id, _assess_id = _bootstrap(c)
    # Seeded but nothing scored: every in-scope row is unscored.
    r = c.post(f"/csf/services/{svc_id}/playbook/export", headers=h)
    assert r.status_code == 409, r.text
    msg = r.json()["error"]["message"]
    assert "unscored" in msg.lower()
    # Count reflects the seeded in-scope rows (one full tier = 106 subcategories).
    assert "of 106" in msg, msg


@pytest.mark.unit
def test_export_gate_blocks_when_not_approved(app_ctx) -> None:
    c, _cid = app_ctx
    h, svc_id, _assess_id = _bootstrap(c)
    _score_all(c, h, svc_id)
    # Every in-scope row scored, but the assessment is still a draft.
    r = c.post(f"/csf/services/{svc_id}/playbook/export", headers=h)
    assert r.status_code == 409, r.text
    assert "not approved" in r.json()["error"]["message"].lower()


@pytest.mark.unit
def test_export_succeeds_when_scored_and_approved(app_ctx) -> None:
    c, cid = app_ctx
    h, svc_id, assess_id = _bootstrap(c)
    _score_all(c, h, svc_id)
    c.post(f"/csf/assessments/{assess_id}/approve", headers=h)

    ex = c.post(f"/csf/services/{svc_id}/playbook/export", headers=h)
    assert ex.status_code == 200, ex.text
    arts = {a["kind"]: a for a in ex.json()["artifacts"]}
    assert "xlsx" in arts

    dl = c.get(
        f"/artifacts/{arts['xlsx']['artifact_id']}/download",
        headers={**h, "X-Client-Id": cid},
    )
    assert dl.status_code == 200
    wb = load_workbook(io.BytesIO(dl.content))
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                assert cell != "Unscored", f"unexpected Unscored cell in sheet {ws.title}"
